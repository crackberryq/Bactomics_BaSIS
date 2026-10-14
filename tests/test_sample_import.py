import csv
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from basis.hybas import enrich_sample
from basis.inputs import load_samples
from basis.sample_import import import_samples
from basis.util import sha256_file

try:
    import openpyxl
except ImportError:
    openpyxl = None


class SampleImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.root = self.base / "inputs"
        self.root.mkdir()
        self.output = self.base / "study"

    def annotations(self, sid="S01", layout="annotation", genbank=False):
        folder = self.root / sid / "annotation" if layout == "annotation" else self.root / sid if layout == "sample" else self.root
        folder.mkdir(parents=True, exist_ok=True)
        contents = {"assembly": (".fna", ">contig\nATGAAATAA\n"),
                    "proteins": (".faa", ">protein\nMK\n"),
                    "gff": (".gff", "##gff-version 3\ncontig\tfixture\tCDS\t1\t9\t.\t+\t0\tID=protein\n")}
        if genbank:
            contents["genbank"] = (".gbff", "fixture GenBank; import checks file identity only\n")
        result = {}
        for key, (suffix, content) in contents.items():
            path = folder / (sid + suffix)
            path.write_text(content)
            result[key] = str(path)
        return result

    def sheet(self, headers=("sample_id",), rows=(("S01",),), suffix=".tsv"):
        path = self.base / ("samples" + suffix)
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle, delimiter="," if suffix == ".csv" else "\t")
            writer.writerow(headers)
            writer.writerows(rows)
        return path

    def workbook(self, rows):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Samples"
        for row in rows:
            worksheet.append(row)
        return workbook

    def save_workbook(self, workbook):
        path = self.base / "samples.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def test_csv_names_discovery_records_provenance_and_source_snapshot(self):
        self.annotations(genbank=True)
        source = self.sheet(suffix=".csv")
        result = import_samples(source, self.output, root=self.root)
        self.assertEqual(result["sample_count"], 1)
        self.assertEqual(result["input_modes"], {"auto": 1})
        audit = json.loads(Path(result["audit"]).read_text())
        self.assertEqual(audit["source"]["sha256"], sha256_file(source))
        self.assertEqual(Path(result["source_sheet"]).read_bytes(), source.read_bytes())
        self.assertEqual(Path(result["source_sheet"]).stat().st_mode & 0o222, 0)
        self.assertEqual(audit["samples"][0]["source_row"], 2)
        self.assertEqual(audit["samples"][0]["discovery_rule"], "root/sample/annotation/sample")
        self.assertIn("genbank", audit["samples"][0]["inputs"])
        self.assertEqual(audit["manifest_sha256"], sha256_file(result["samples"]))

    def test_all_exact_layouts_and_external_mode(self):
        for sid, layout in (("S01", "annotation"), ("S02", "sample"), ("S03", "flat")):
            self.annotations(sid, layout)
        source = self.sheet(rows=(("S01",), ("S02",), ("S03",)))
        result = import_samples(source, self.output, root=self.root, input_mode="external")
        loaded = load_samples(result["samples"])
        self.assertEqual(len(loaded), 3)
        self.assertTrue(all(row["input_mode"] == "external" for row in loaded))
        self.assertTrue(all("hybas_context" not in row for row in loaded))

    def test_explicit_relative_paths_do_not_use_discovery_root(self):
        paths = self.annotations(layout="flat")
        source = self.sheet(headers=("sample_id", "assembly", "proteins", "gff", "taxon", "taxon_phylum"),
                            rows=(("S01", *(str(Path(paths[key]).relative_to(self.base)) for key in ("assembly", "proteins", "gff")), "Example genus", "Example phylum"),))
        result = import_samples(source, self.output, root=self.base, input_mode="external")
        sample = load_samples(result["samples"])[0]
        self.assertEqual(sample["assembly"], paths["assembly"])
        self.assertEqual(sample["taxon_status"], "declared")
        self.assertEqual(sample["taxon_source"], "sample_sheet")
        audit = json.loads(Path(result["audit"]).read_text())
        self.assertEqual(audit["samples"][0]["taxonomy_basis"], "user_declaration")

    def test_source_mode_per_row_overrides_default(self):
        self.annotations()
        result = import_samples(self.sheet(headers=("sample_id", "input_mode"), rows=(("S01", "external"),)),
                                self.output, root=self.root, input_mode="hybas")
        self.assertEqual(load_samples(result["samples"])[0]["input_mode"], "external")

    def test_config_is_created_and_hashed_in_same_import(self):
        self.annotations()
        result = import_samples(self.sheet(), self.output, root=self.root,
                                config={"samples": "samples.tsv", "outdir": "results"})
        audit = json.loads(Path(result["audit"]).read_text())
        self.assertEqual(audit["config_sha256"], sha256_file(result["config"]))

    def test_config_exclusion_for_absent_sample_fails_before_staging(self):
        self.annotations()
        config = {"samples": "samples.tsv", "sample_module_overrides": {
            "absent": {"pathogenfinder2": {"enabled": False, "reason": "declared exclusion"}}
        }}
        with patch("basis.sample_import.tempfile.mkdtemp") as stage:
            with self.assertRaisesRegex(ValueError, "unknown sample IDs"):
                import_samples(self.sheet(), self.output, root=self.root, config=config)
            stage.assert_not_called()
        self.assertFalse(self.output.exists())

    def test_config_valid_sample_exclusion_is_preserved(self):
        self.annotations()
        import yaml
        config = {"sample_module_overrides": {
            "S01": {"pathogenfinder2": {"enabled": False, "reason": "declared exclusion"}}
        }}
        result = import_samples(self.sheet(), self.output, root=self.root, config=config)
        self.assertEqual(yaml.safe_load(Path(result["config"]).read_text()), config)

    def test_ambiguous_layouts_never_choose_one(self):
        self.annotations()
        self.annotations(layout="flat")
        with self.assertRaisesRegex(ValueError, "ambiguous discovery"):
            import_samples(self.sheet(), self.output, root=self.root)
        self.assertFalse(self.output.exists())

    def test_ambiguous_optional_genbank_is_rejected(self):
        paths = self.annotations(genbank=True)
        Path(paths["genbank"]).with_suffix(".gbk").write_text("another")
        with self.assertRaisesRegex(ValueError, "both .gbk and .gbff"):
            import_samples(self.sheet(), self.output, root=self.root)

    def test_partial_explicit_and_discovered_inputs_require_annotation(self):
        paths = self.annotations()
        with self.assertRaisesRegex(ValueError, "annotation_required"):
            import_samples(self.sheet(headers=("sample_id", "assembly"), rows=(("S01", paths["assembly"]),)), self.output, root=self.root)
        Path(paths["proteins"]).unlink()
        with self.assertRaisesRegex(ValueError, "annotation_required"):
            import_samples(self.sheet(), self.output, root=self.root)
        self.assertFalse(self.output.exists())

    def test_no_root_and_no_matching_files_fail_clearly(self):
        with self.assertRaisesRegex(ValueError, "require --root"):
            import_samples(self.sheet(), self.output)
        with self.assertRaisesRegex(ValueError, "no annotation files"):
            import_samples(self.sheet(), self.output, root=self.root)

    def test_duplicate_casefold_sample_ids_and_invalid_ids_rejected(self):
        for rows in ((("S01",), ("s01",)), (("../bad",),), (("",),)):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                import_samples(self.sheet(rows=rows), self.output, root=self.root)
        self.assertFalse(self.output.exists())

    def test_malformed_headers_and_row_width_fail_before_discovery(self):
        for header, rows in ((("sample_id", "sample_id"), (("S01", "S02"),)),
                             (("sample_id", "typo"), (("S01", "x"),)),
                             (("sample_id", ""), (("S01", ""),)),
                             (("sample_id",), (("S01", "extra"),)),
                             (("sample_id", "assembly"), (("S01",),))):
            with self.subTest(header=header, rows=rows), self.assertRaises(ValueError):
                import_samples(self.sheet(header, rows), self.output, root=self.root)

    def test_shared_required_input_paths_symlinks_and_hardlinks_rejected(self):
        first = self.annotations("S01", "flat")
        second = self.annotations("S02", "flat")
        for method in ("path", "symlink", "hardlink"):
            other = dict(second)
            if method == "path":
                other["proteins"] = first["proteins"]
            else:
                link = self.base / ("alias-" + method)
                getattr(link, "symlink_to" if method == "symlink" else "hardlink_to")(first["proteins"])
                other["proteins"] = str(link)
            rows = [(sid, *(paths[key] for key in ("assembly", "proteins", "gff")))
                    for sid, paths in (("S01", first), ("S02", other))]
            source = self.sheet(("sample_id", "assembly", "proteins", "gff"), rows)
            with self.subTest(method=method), self.assertRaisesRegex(ValueError, "shared required input"):
                import_samples(source, self.output, input_mode="external")

    def test_missing_explicit_path_is_error_without_output(self):
        paths = self.annotations()
        Path(paths["gff"]).unlink()
        source = self.sheet(("sample_id", "assembly", "proteins", "gff"),
                            (("S01", *(paths[key] for key in ("assembly", "proteins", "gff"))),))
        with self.assertRaisesRegex(ValueError, "file does not exist"):
            import_samples(source, self.output)
        self.assertFalse(self.output.exists())

    def test_existing_output_even_empty_is_preserved(self):
        self.annotations()
        self.output.mkdir()
        with self.assertRaises(FileExistsError):
            import_samples(self.sheet(), self.output, root=self.root)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_concurrently_created_empty_output_is_not_replaced(self):
        self.annotations()
        original_copy = __import__("shutil").copyfile
        def concurrent_copy(src, dst):
            result = original_copy(src, dst)
            self.output.mkdir()
            return result
        with patch("basis.sample_import.shutil.copyfile", side_effect=concurrent_copy):
            with self.assertRaises(FileExistsError):
                import_samples(self.sheet(), self.output, root=self.root)
        self.assertTrue(self.output.is_dir())
        self.assertEqual(list(self.output.iterdir()), [])

    def test_source_mutation_detected_before_commit(self):
        self.annotations()
        source = self.sheet()
        original_copy = __import__("shutil").copyfile
        def changing_copy(src, dst):
            result = original_copy(src, dst)
            source.write_text("sample_id\nchanged\n")
            return result
        with patch("basis.sample_import.shutil.copyfile", side_effect=changing_copy):
            with self.assertRaisesRegex(ValueError, "Source sheet changed"):
                import_samples(source, self.output, root=self.root)
        self.assertFalse(self.output.exists())
        self.assertFalse(list(self.base.glob(".basis-import-*")))

    def test_input_mutation_detected_before_commit(self):
        paths = self.annotations()
        original_copy = __import__("shutil").copyfile
        def changing_copy(src, dst):
            result = original_copy(src, dst)
            Path(paths["assembly"]).write_text(">changed\nA\n")
            return result
        with patch("basis.sample_import.shutil.copyfile", side_effect=changing_copy):
            with self.assertRaisesRegex(ValueError, "Input changed"):
                import_samples(self.sheet(), self.output, root=self.root)
        self.assertFalse(self.output.exists())

    def test_commit_error_rolls_back_new_output(self):
        self.annotations()
        original_rename = Path.rename
        def failing_rename(path, target):
            if path.name == "samples.tsv":
                raise OSError("synthetic disk error")
            return original_rename(path, target)
        with patch.object(Path, "rename", failing_rename):
            with self.assertRaisesRegex(OSError, "synthetic disk error"):
                import_samples(self.sheet(), self.output, root=self.root)
        self.assertFalse(self.output.exists())
        self.assertFalse(list(self.base.glob(".basis-import-*")))

    def test_unsupported_legacy_excel_has_save_as_guidance(self):
        source = self.base / "old.xls"
        source.write_bytes(b"legacy")
        with self.assertRaisesRegex(ValueError, "save the workbook as .xlsx"):
            import_samples(source, self.output)

    def test_sheet_selector_on_delimited_file_is_error(self):
        with self.assertRaisesRegex(ValueError, "only applicable"):
            import_samples(self.sheet(), self.output, sheet_name="Samples")

    def test_missing_excel_extra_has_actionable_message(self):
        source = self.base / "samples.xlsx"
        source.write_bytes(b"contents are not opened before dependency check")
        original_import = __import__
        def missing_excel(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("synthetic missing dependency")
            return original_import(name, *args, **kwargs)
        with patch("builtins.__import__", side_effect=missing_excel):
            with self.assertRaisesRegex(ValueError, "bactromics-basis\\[excel\\]"):
                import_samples(source, self.output)

    @unittest.skipIf(openpyxl is None, "openpyxl extra not installed")
    def test_xlsx_text_identifiers_preserve_leading_zeros(self):
        self.annotations("001")
        source = self.save_workbook(self.workbook([["sample_id"], ["001"]]))
        result = import_samples(source, self.output, root=self.root, input_mode="external")
        self.assertEqual(load_samples(result["samples"])[0]["sample_id"], "001")
        self.assertEqual(json.loads(Path(result["audit"]).read_text())["source"]["sheet"], "Samples")

    @unittest.skipIf(openpyxl is None, "openpyxl extra not installed")
    def test_xlsx_numeric_identifier_formula_and_error_are_rejected(self):
        for value in (1, "=1+1", "#DIV/0!"):
            source = self.save_workbook(self.workbook([["sample_id"], [value]]))
            with self.subTest(value=value), self.assertRaises(ValueError):
                import_samples(source, self.output, root=self.root)

    @unittest.skipIf(openpyxl is None, "openpyxl extra not installed")
    def test_xlsx_multiple_sheets_require_explicit_selection(self):
        self.annotations()
        workbook = self.workbook([["sample_id"], ["S01"]])
        workbook.create_sheet("Other")
        source = self.save_workbook(workbook)
        with self.assertRaisesRegex(ValueError, "multiple worksheets"):
            import_samples(source, self.output, root=self.root)
        result = import_samples(source, self.output, root=self.root, sheet_name="Samples")
        self.assertEqual(result["sample_count"], 1)

    @unittest.skipIf(openpyxl is None, "openpyxl extra not installed")
    def test_xlsx_merged_cells_are_rejected(self):
        workbook = self.workbook([["sample_id"], ["S01"]])
        workbook.active.merge_cells("A1:B1")
        source = self.save_workbook(workbook)
        with self.assertRaisesRegex(ValueError, "Merged cells"):
            import_samples(source, self.output, root=self.root)


class InputModeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        annotation = self.base / "S01" / "annotation"
        annotation.mkdir(parents=True)
        self.sample = {"sample_id": "S01"}
        for key, suffix in (("assembly", ".fna"), ("proteins", ".faa"), ("gff", ".gff")):
            path = annotation / ("S01" + suffix)
            path.write_text("fixture")
            self.sample[key] = str(path)
        self.manifest = self.base / "samples.tsv"

    def write_manifest(self, headers, rows):
        with self.manifest.open("w", newline="") as handle:
            writer = csv.writer(handle, delimiter="\t")
            writer.writerow(headers)
            writer.writerows(rows)

    def test_external_does_not_read_matching_hybas_metadata(self):
        metadata = self.base / "S01" / "metadata"
        metadata.mkdir()
        (metadata / "run_summary.json").write_text("this is deliberately invalid metadata")
        external = enrich_sample({**self.sample, "input_mode": "external", "taxon": "Declared species"})
        self.assertNotIn("hybas_context", external)
        self.assertEqual(external["taxon_status"], "declared")
        with self.assertRaises(ValueError):
            enrich_sample(self.sample)

    def test_explicit_external_and_hybas_root_conflict(self):
        with self.assertRaisesRegex(ValueError, "conflicts"):
            enrich_sample({**self.sample, "input_mode": "external", "hybas_dir": str(self.base / "S01")})

    def test_unknown_taxonomy_stays_unknown_in_external_and_hybas(self):
        for mode in ("external", "hybas"):
            result = enrich_sample({**self.sample, "input_mode": mode})
            self.assertNotIn("taxon", result)
        self.assertEqual(enrich_sample(self.sample)["input_mode"], "auto")

    def test_explicit_hybas_mode_requires_matching_layout(self):
        with self.assertRaisesRegex(ValueError, "requires a matching"):
            enrich_sample({**self.sample, "gff": str(self.base / "other.gff"), "input_mode": "hybas"})

    def test_invalid_mode_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Invalid input_mode"):
            enrich_sample({**self.sample, "input_mode": "guess"})

    def test_loader_strict_column_counts_headers_and_casefold_ids(self):
        headers = ("sample_id", "assembly", "proteins", "gff", "input_mode")
        row = ("S01", self.sample["assembly"], self.sample["proteins"], self.sample["gff"], "external")
        for current_headers, rows in ((headers + ("sample_id",), (row + ("S01",),)),
                                      (headers, (row + ("extra",),)),
                                      (headers, (row[:-1],)),
                                      (headers, (row, ("s01", *row[1:])))):
            self.write_manifest(current_headers, rows)
            with self.subTest(headers=current_headers, rows=rows), self.assertRaises(ValueError):
                load_samples(self.manifest)
        self.write_manifest(headers, (row,))
        self.assertEqual(load_samples(self.manifest)[0]["input_mode"], "external")

    def test_loader_converts_malformed_quoted_tsv_to_value_error(self):
        for content in ('"sample_id\tassembly\tproteins\tgff\n',
                        'sample_id\tassembly\tproteins\tgff\nS01\t"unterminated\n'):
            self.manifest.write_text(content)
            with self.subTest(content=content), self.assertRaisesRegex(ValueError, "malformed quoted TSV"):
                load_samples(self.manifest)


if __name__ == "__main__":
    unittest.main()
