"""Import sample manifests without running annotation or biological analyses."""
from __future__ import annotations

from collections import Counter
import csv
import json
from pathlib import Path
import re
import shutil
import tempfile
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from .inputs import SAMPLE_ID, normalize_input_mode
from .progress import emit
from .util import now, resolve_path, sha256_file

FIELDS = (
    "sample_id", "assembly", "proteins", "gff", "genbank", "taxon",
    "taxon_rank", "taxon_source", "amrfinder_organism", "taxon_phylum",
    "phenotypes", "hybas_dir", "input_mode",
)
REQUIRED_INPUTS = {"assembly": ".fna", "proteins": ".faa", "gff": ".gff"}


def _cell_text(value, label):
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError(f"{label}: enter text, not numeric/date/boolean cells; sample IDs must preserve leading zeros")
    if any(char in value for char in ("\x00", "\r", "\n", "\t")):
        raise ValueError(f"{label}: control characters are not allowed in manifest cells")
    return value.strip()


def _read_sheet(path, sheet_name):
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise ValueError("Legacy .xls is unsupported; save the workbook as .xlsx, CSV or TSV")
    if suffix not in {".xlsx", ".csv", ".tsv"}:
        raise ValueError("Sample import accepts .xlsx, .csv or .tsv files")
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("Excel import requires openpyxl; install the bactromics-basis[excel] extra, or export CSV/TSV") from exc
        # data_only=False is essential: formula cells must not silently become
        # cached values or blanks. External workbook links are not followed.
        workbook = openpyxl.load_workbook(path, data_only=False, keep_links=False)
        try:
            if sheet_name is None:
                if len(workbook.sheetnames) != 1:
                    raise ValueError("Workbook has multiple worksheets; select one with --sheet")
                sheet_name = workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet {sheet_name!r} does not exist")
            worksheet = workbook[sheet_name]
            if worksheet.merged_cells.ranges:
                raise ValueError("Merged cells are not supported; use a rectangular sample table")
            if worksheet.max_row > 100001 or worksheet.max_column > 100:
                raise ValueError("Worksheet exceeds 100,000 sample rows or 100 columns; remove distant formatted cells")
            rows = []
            for number, cells in enumerate(worksheet.iter_rows(), 1):
                values = []
                for cell in cells:
                    if cell.data_type in {"f", "e"}:
                        raise ValueError(f"{sheet_name}!{cell.coordinate}: formulas and Excel errors are not allowed")
                    values.append(_cell_text(cell.value, f"{sheet_name}!{cell.coordinate}"))
                rows.append((number, values))
        finally:
            workbook.close()
        # Ignore physically empty trailing columns. Never trim actual cell data.
        width = max((i + 1 for _, row in rows for i, value in enumerate(row) if value), default=0)
        rows = [(number, row[:width]) for number, row in rows]
        return rows, sheet_name
    if sheet_name is not None:
        raise ValueError("--sheet is only applicable to .xlsx workbooks")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter="," if suffix == ".csv" else "\t", strict=True)
        try:
            return [(number, [_cell_text(value, f"row {number}") for value in row])
                    for number, row in enumerate(reader, 1)], None
        except csv.Error as exc:
            raise ValueError(f"Malformed delimited sample sheet: {exc}") from exc


def _table(rows):
    if not rows or not any(rows[0][1]):
        raise ValueError("Sample sheet must begin with a header row")
    headers = rows[0][1]
    if (any(not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", key) for key in headers)
            or len({key.casefold() for key in headers}) != len(headers)):
        raise ValueError("Sample sheet has empty, malformed or duplicate column headers")
    unknown = set(headers) - set(FIELDS)
    if unknown:
        raise ValueError("Unsupported sample columns: " + ", ".join(sorted(unknown)))
    if "sample_id" not in headers:
        raise ValueError("Sample sheet requires a sample_id column")
    result = []
    seen = set()
    for number, cells in rows[1:]:
        if not any(cells):
            continue
        if len(cells) != len(headers):
            raise ValueError(f"Sample sheet row {number}: expected {len(headers)} columns, found {len(cells)}")
        row = dict(zip(headers, cells))
        sid = row["sample_id"]
        if not SAMPLE_ID.fullmatch(sid) or sid in {".", ".."}:
            raise ValueError(f"Row {number}: invalid sample ID {sid!r}")
        if sid.casefold() in seen:
            raise ValueError(f"Row {number}: duplicate sample ID {sid!r} (case-insensitive portability check)")
        seen.add(sid.casefold())
        result.append((number, row))
    if not result:
        raise ValueError("Sample sheet contains no samples")
    return result


def _discover(root, sid):
    layouts = (
        ("root/sample/annotation/sample", root / sid / "annotation" / sid),
        ("root/sample/sample", root / sid / sid),
        ("root/sample", root / sid),
    )
    found = []
    for rule, stem in layouts:
        paths = {key: Path(str(stem) + suffix) for key, suffix in REQUIRED_INPUTS.items()}
        genbanks = [Path(str(stem) + suffix) for suffix in (".gbk", ".gbff")]
        if any(path.exists() for path in [*paths.values(), *genbanks]):
            found.append((rule, paths, genbanks))
    if len(found) > 1:
        raise ValueError(f"{sid}: ambiguous discovery: files occur in multiple supported layouts; supply explicit paths")
    if not found:
        raise ValueError(f"{sid}: no annotation files found in the three supported exact layouts under {root}")
    rule, paths, genbanks = found[0]
    missing = [key for key, path in paths.items() if not path.is_file()]
    if missing:
        raise ValueError(f"{sid}: annotation_required: missing {', '.join(missing)}; provide matching assembly, protein FASTA and GFF files; BaSIS does not run annotation")
    existing_gb = [path for path in genbanks if path.exists()]
    if len(existing_gb) > 1:
        raise ValueError(f"{sid}: both .gbk and .gbff exist; supply explicit paths to choose the annotation set")
    if existing_gb:
        paths["genbank"] = existing_gb[0]
    return {key: str(path.resolve()) for key, path in paths.items()}, rule


def _mode_structure(sample):
    """Check declared source layout, leaving sequence/metadata checks to validate."""
    mode, sid = sample["input_mode"], sample["sample_id"]
    explicit = sample.get("hybas_dir")
    if mode == "external" and explicit:
        raise ValueError(f"{sid}: input_mode external conflicts with hybas_dir")
    if mode != "hybas" and not explicit:
        return
    gff = Path(sample["gff"])
    root = Path(explicit) if explicit else gff.parent.parent
    if (root.name != sid or not root.is_dir()
            or any(Path(sample[key]) != (root / "annotation" / (sid + suffix)).resolve()
                   for key, suffix in REQUIRED_INPUTS.items())):
        raise ValueError(f"{sid}: hybas input requires one matching isolate/annotation set and identifiable root")
    if sample.get("genbank") and Path(sample["genbank"]) not in {
        (root / "annotation" / (sid + suffix)).resolve() for suffix in (".gbk", ".gbff")
    }:
        raise ValueError(f"{sid}: HybAs GenBank must belong to the same annotation set")
    sample["hybas_dir"] = str(root)


def _resolve_row(number, row, source, selected_sheet, mode, discovery_root):
    """Resolve one row without writes; leave cross-row file identity to the caller."""
    sid = row["sample_id"]
    sample = {key: row.get(key, "") for key in FIELDS}
    errors, inputs, identities = [], {}, []
    rule = None
    try:
        sample["input_mode"] = normalize_input_mode(row.get("input_mode") or mode)
    except ValueError as exc:
        errors.append(f"{sid}: {exc}")
    supplied = [key for key in REQUIRED_INPUTS if sample[key]]
    if supplied:
        rule = "explicit_paths_relative_to_source_sheet"
        if len(supplied) != len(REQUIRED_INPUTS):
            errors.append(f"{sid}: annotation_required: supply all assembly, proteins and gff paths together, or only sample_id with --root")
    elif discovery_root is None:
        errors.append(f"{sid}: names-only rows require --root, or provide assembly, proteins and gff paths")
    else:
        try:
            discovered, rule = _discover(discovery_root, sid)
            for key, value in discovered.items():
                if key != "genbank" or not sample["genbank"]:
                    sample[key] = value
        except (ValueError, OSError) as exc:
            errors.append(str(exc))
    for key in (*REQUIRED_INPUTS, "genbank", "phenotypes"):
        if not sample.get(key):
            continue
        try:
            sample[key] = resolve_path(sample[key], source.parent)
            path = Path(sample[key])
            if not path.is_file():
                errors.append(f"{sid}: {key} file does not exist: {path}")
                continue
            stat = path.stat()
            if key in REQUIRED_INPUTS:
                identities.append((key, (str(path).casefold(), (stat.st_dev, stat.st_ino))))
            inputs[key] = {"path": str(path), "sha256": sha256_file(path), "bytes": stat.st_size}
        except (ValueError, OSError) as exc:
            errors.append(f"{sid}: cannot read {key}: {exc}")
    if sample["hybas_dir"]:
        try:
            sample["hybas_dir"] = resolve_path(sample["hybas_dir"], source.parent)
            if not Path(sample["hybas_dir"]).is_dir():
                errors.append(f"{sid}: hybas_dir does not exist")
        except (ValueError, OSError) as exc:
            errors.append(f"{sid}: cannot resolve hybas_dir: {exc}")
    if not errors:
        try:
            _mode_structure(sample)
        except (ValueError, OSError) as exc:
            errors.append(str(exc))
    if sample["taxon"] and not sample["taxon_source"]:
        sample["taxon_source"] = "sample_sheet"
    if sample["input_mode"] == "external" and sample["taxon_source"] == "hybas_kraken_operational_target":
        errors.append(f"{sid}: external input cannot claim an imported HybAs taxonomy source")
    audit = {"sample_id": sid, "source_row": number, "source_sheet": selected_sheet,
             "discovery_rule": rule, "input_mode": sample["input_mode"],
             "taxonomy_basis": "user_declaration" if sample["taxon"] or sample["taxon_phylum"] else "not_recorded",
             "inputs": inputs}
    return sample, audit, errors, identities


def _prepare_samples(sheet, root, sheet_name, input_mode, config):
    """Shared read-only import validation, including all readable row failures."""
    source = Path(sheet).expanduser().resolve(strict=True)
    if not source.is_file():
        raise ValueError(f"Source sheet is not a file: {source}")
    if source.stat().st_size > 50 * 1024 * 1024:
        raise ValueError("Sample sheet exceeds 50 MiB")
    mode = normalize_input_mode(input_mode)
    discovery_root = Path(root).expanduser().resolve(strict=True) if root is not None else None
    if discovery_root is not None and not discovery_root.is_dir():
        raise ValueError("Discovery root must be a directory")
    source_hash = sha256_file(source)
    cells, selected_sheet = _read_sheet(source, sheet_name)
    rows = _table(cells)
    samples, audits, row_errors, paths_seen = [], [], [], {}
    shared_errors = {}
    for number, row in rows:
        emit(f"Sample import {row['sample_id']} (row {number}): checking input files")
        sample, audit, errors, identities = _resolve_row(number, row, source, selected_sheet, mode, discovery_root)
        index = len(samples)
        samples.append(sample)
        audits.append(audit)
        row_errors.append(errors)
        for key, identifiers in identities:
            # Keep one representative for each canonical path and inode. Every
            # alias participant is blocked, but a large repeated-file group
            # does not produce all pairwise combinations of diagnostics.
            previous = {paths_seen[identity] for identity in identifiers if identity in paths_seen}
            for previous_index, previous_key in sorted(previous):
                previous_id = samples[previous_index]["sample_id"]
                message = (f"shared required input path/file: {previous_id}:{previous_key} "
                           f"and {sample['sample_id']}:{key}")
                shared_errors.setdefault(previous_index, message)
                shared_errors.setdefault(index, message)
            for identity in identifiers:
                paths_seen.setdefault(identity, (index, key))
    # One representative shared-file explanation per affected row is enough to
    # reject the transaction. Diagnostics and identity storage remain linear.
    for index, message in shared_errors.items():
        row_errors[index].append(message)
    for sample, errors in zip(samples, row_errors):
        state = "blocked" if errors else "ready"
        emit(f"Sample import {sample['sample_id']}: {state}")
    errors = []
    if config is not None:
        try:
            if not isinstance(config, dict):
                raise ValueError("config must be a caller-prepared dictionary")
            from .sample_policy import validate_sample_overrides
            validate_sample_overrides(config, samples)
        except ValueError as exc:
            errors.append(str(exc))
    if sha256_file(source) != source_hash:
        errors.append("Source sheet changed during preview/import; run again with an unchanged sheet")
    return {"source": source, "source_hash": source_hash, "selected_sheet": selected_sheet,
            "discovery_root": discovery_root, "samples": samples, "audits": audits,
            "row_errors": row_errors, "errors": errors}


def preview_samples(sheet, root=None, sheet_name=None, input_mode="auto", config=None):
    """Read-only import preview. No study, staging directory or other files are written.

    Ready/blocked sample counts describe row input checks. Top-level structural
    or configuration errors can block the transaction independently. Sequence
    consistency, resource readiness and biological suitability are not checked.
    """
    from .input_context import describe_input_context
    result = {"schema_version": 1, "operation": "sample_manifest_preview",
              "status": "blocked", "sample_count": 0, "ready_samples": 0,
              "blocked_samples": 0, "samples": [], "errors": [],
              "sequence_consistency_checked": False, "resources_checked": False,
              "native_analyses_run": False}
    try:
        prepared = _prepare_samples(sheet, root, sheet_name, input_mode, config)
    except (ValueError, OSError, BadZipFile, ParseError) as exc:
        result["errors"].append(str(exc))
        return result
    result["source"] = {"path": str(prepared["source"]), "sha256": prepared["source_hash"],
                        "sheet": prepared["selected_sheet"]}
    result["root"] = str(prepared["discovery_root"]) if prepared["discovery_root"] else None
    result["errors"] = prepared["errors"]
    for sample, audit, errors in zip(prepared["samples"], prepared["audits"], prepared["row_errors"]):
        result["samples"].append({
            "sample_id": sample["sample_id"], "source_row": audit["source_row"],
            "status": "blocked" if errors else "ready", "input_mode": sample["input_mode"],
            "input_context": describe_input_context(sample, resolved=False),
            "discovery_rule": audit["discovery_rule"],
            "resolved_inputs": {key: value["path"] for key, value in audit["inputs"].items()},
            "errors": errors})
    result["sample_count"] = len(result["samples"])
    result["ready_samples"] = sum(sample["status"] == "ready" for sample in result["samples"])
    result["blocked_samples"] = result["sample_count"] - result["ready_samples"]
    if not result["errors"] and not result["blocked_samples"]:
        result["status"] = "ready_for_import"
    return result


def import_samples(sheet, outdir, root=None, sheet_name=None, input_mode="auto", config=None):
    """Create a new TSV manifest and provenance audit; never run native tools.

    ``config`` is an optional caller-prepared dictionary saved as config.yaml in
    the same transaction. All explicit input paths are relative to ``sheet``.
    The output directory must not exist, even if empty.
    """
    output = Path(outdir).expanduser().absolute()
    if output.exists() or output.is_symlink():
        raise FileExistsError(f"Import output already exists: {output}")
    prepared = _prepare_samples(sheet, root, sheet_name, input_mode, config)
    problems = [*prepared["errors"], *(error for row in prepared["row_errors"] for error in row)]
    if problems:
        raise ValueError("\n".join(dict.fromkeys(problems)))
    source = prepared["source"]
    source_hash = prepared["source_hash"]
    selected_sheet = prepared["selected_sheet"]
    discovery_root = prepared["discovery_root"]
    samples, audits = prepared["samples"], prepared["audits"]

    # All checks precede staging. Recheck the original sheet and input bytes just
    # before commit so the audit never deliberately records a stale snapshot.
    output.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".basis-import-", dir=output.parent))
    created_output = False
    try:
        copied_name = "source-sheet" + source.suffix.lower()
        shutil.copyfile(source, staging / copied_name)
        if sha256_file(staging / copied_name) != source_hash or sha256_file(source) != source_hash:
            raise ValueError("Source sheet changed during import; no output committed")
        (staging / copied_name).chmod(0o444)
        with (staging / "samples.tsv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS, delimiter="\t", lineterminator="\n")
            writer.writeheader()
            writer.writerows(samples)
        audit = {"schema_version": 1, "created_at": now(), "operation": "sample_manifest_import",
                 "source": {"path": str(source), "snapshot": copied_name, "sha256": source_hash,
                            "sheet": selected_sheet}, "root": str(discovery_root) if discovery_root else None,
                 "sample_count": len(samples), "samples": audits,
                 "manifest_sha256": sha256_file(staging / "samples.tsv"),
                 "limitations": ["Import checks file identity and table structure; run basis validate/doctor for sequence consistency.",
                                 "Taxonomy values are declarations; this import performs no taxonomic classification or biological validation."]}
        if config is not None:
            import yaml
            (staging / "config.yaml").write_text(yaml.safe_dump(config, sort_keys=False), encoding="utf-8")
            audit["config_sha256"] = sha256_file(staging / "config.yaml")
        (staging / "import-audit.json").write_text(json.dumps(audit, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        for sample_audit in audits:
            for record in sample_audit["inputs"].values():
                if sha256_file(record["path"]) != record["sha256"]:
                    raise ValueError(f"Input changed during import: {record['path']}; no output committed")
        if sha256_file(source) != source_hash:
            raise ValueError("Source sheet changed during import; no output committed")
        # Exclusive creation guarantees an existing directory is never replaced,
        # including an empty one created concurrently. Any commit error removes
        # only the directory this call just created.
        output.mkdir(exist_ok=False)
        created_output = True
        for artifact in staging.iterdir():
            artifact.rename(output / artifact.name)
    except Exception:
        if created_output:
            shutil.rmtree(output)
        raise
    finally:
        shutil.rmtree(staging)
    summary = {"status": "imported", "sample_count": len(samples),
               "samples": str(output / "samples.tsv"), "audit": str(output / "import-audit.json"),
               "source_sheet": str(output / copied_name),
               "input_modes": dict(Counter(row["input_mode"] for row in samples)),
               "note": "No analyses run. Validate inputs and review module decisions before execution."}
    if config is not None:
        summary["config"] = str(output / "config.yaml")
    return summary
