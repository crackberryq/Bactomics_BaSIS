"""Read-only workflow configuration; no tools or downloads during DAG creation."""
from __future__ import annotations
import copy
import csv
import json
import os
from pathlib import Path
import re

from basis.config import ALL_MODULES, PATH_KEYS, default_config, merge_dict
from basis.hybas import enrich_sample
from basis.sample_policy import plan_sample_modules, validate_sample_overrides
from .runtime import _positive_seconds

RECIPES = {"amrfinder": "core-amr", "rgi": "rgi", "vfdb": "diamond", "bacmet": "diamond",
           "antismash": "antismash", "genomad": "context", "integronfinder": "context",
           "mobsuite": "mobsuite", "isescan": "isescan", "icescreen": "icescreen",
           "crisprcasfinder": "crisprcasfinder", "minced": "minced", "mlst": "mlst",
           "pathogenfinder2": "pathogenfinder2", "deepsafe": "deepsafe"}
PROVIDERS = {module: module for module in ALL_MODULES if module != "pathogenfinder2"}
PROVIDERS["rgi"] = "card"
PROVIDER_OWNER = {value: key for key, value in PROVIDERS.items()}
PROVIDER_OWNER.update(prott5="pathogenfinder2", prosite="deepsafe", deep_catalogue="deepsafe")


def _bool(value, label):
    if not isinstance(value, bool):
        raise ValueError(label + " must be true or false")
    return value


def _positive(value, label):
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError(label + " must be a positive integer")
    return value


def _path(value, base):
    path = Path(os.path.expandvars(str(value))).expanduser()
    return str((base / path).resolve())


def _sheet(path):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(book.active.iter_rows(values_only=True))
        finally:
            book.close()
    else:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle, delimiter="," if path.suffix.lower() == ".csv" else "\t"))
    if not rows:
        raise ValueError("Sample sheet is empty")
    headers = [str(value or "").strip() for value in rows[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("Sample sheet columns must be nonempty and unique")
    result = []
    for number, cells in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        if len(cells) != len(headers):
            raise ValueError(f"Sample sheet row {number}: wrong number of columns")
        result.append({key: str(value or "").strip() for key, value in zip(headers, cells)})
    return result


class SafetyContext:
    def __init__(self, supplied, root):
        self.root = Path(root).resolve()
        self.supplied = copy.deepcopy(supplied)
        self.base = Path(_path(supplied.get("bactomics_root", ".."), self.root))
        self.threads = _positive(supplied.get("threads", 4), "threads")
        self.memory_mb = _positive(supplied.get("memory_mb", 20000), "memory_mb")
        self.extend_enabled = _bool(supplied.get("extend", True), "extend")
        deep = supplied.get("deep", {})
        if isinstance(deep, bool):
            deep = {"enabled": deep}
        if not isinstance(deep, dict):
            raise ValueError("deep must be a mapping")
        self.deep_enabled = _bool(deep.get("enabled", False), "deep.enabled")
        cfg = default_config()
        for row in cfg["modules"].values():
            row.update(enabled=True, threads=self.threads)
        cfg["deep"].update(enabled=self.deep_enabled, threads=self.threads,
            prosite_profiles=True, auto_phylogeny=True, family_placement=True)
        cfg = merge_dict(cfg, {key: supplied[key] for key in
            ("project", "modules", "sample_module_overrides", "execution") if key in supplied})
        cfg["deep"] = merge_dict(cfg["deep"], deep)
        from basis.aware import resolve_policy
        resolve_policy(cfg["deep"])
        cfg["execution"].update(cores=self.threads, memory_mb=self.memory_mb)
        cfg["version"] = "0.4.0a3"
        cfg["engine"] = "snakemake"
        cfg["options"] = {"deep": self.deep_enabled, "extend": self.extend_enabled}
        if set(cfg["modules"]) - set(ALL_MODULES):
            raise ValueError("Unknown module names: " + str(sorted(set(cfg["modules"]) - set(ALL_MODULES))))
        for name, section in list(cfg["modules"].items()) + [("deep", cfg["deep"])]:
            if "enabled" in section:
                _bool(section["enabled"], name + ".enabled")
            if section.get("command_prefix"):
                raise ValueError("Use Snakemake conda environments; command_prefix is unsupported")
            if section.get("mode", "native") != "native":
                raise ValueError("This standalone workflow runs native tools; imports need a separate reviewed report")
            for key in PATH_KEYS:
                if section.get(key):
                    section[key] = _path(section[key], self.root)
            for key in ("threads", "memory_mb"):
                _positive(section.get(key, 1), name + "." + key)
            for key in ("timeout_seconds", "heartbeat_seconds"):
                if key in section:
                    _positive_seconds(section[key], name + "." + key)
            section["command_prefix"] = []
        cfg["modules"]["pathogenfinder2"]["enabled"] &= self.extend_enabled
        ccf = cfg["modules"]["crisprcasfinder"]
        if ccf.get("execution_scope", "assembly") != "assembly":
            raise ValueError("crisprcasfinder.execution_scope must be assembly")
        ccf["execution_scope"] = "assembly"
        # Older overlays may contain a scheduler batch size. The active rule now
        # runs the complete assembly; this setting no longer affects execution.
        ccf.pop("batch_size", None)
        cfg["deep"].setdefault("executables", {})["iqtree"] = "iqtree2"
        resources = supplied.get("resources", {})
        if not isinstance(resources, dict) or resources.get("policy", "auto") not in {"auto", "manual"}:
            raise ValueError("resources.policy must be auto or manual")
        self.resource_root = _path(resources.get("directory", "resources"), self.root)
        self.resource_policy = resources.get("policy", "auto")
        self.config = cfg
        self.samples = self._samples(supplied)
        validate_sample_overrides(cfg, list(self.samples.values()))
        self.decisions = {isolate: plan_sample_modules(cfg, sample, self.extend_enabled)
                          for isolate, sample in self.samples.items()}
        self._specs = self._resource_specs()
        self._check_resource_links()

    def _check_resource_links(self):
        """Detect missing undeclared payload links without changing the DAG's files."""
        for name, spec in self._specs.items():
            if spec.get("mode") == "external":
                continue
            receipt = Path(self.resource_root) / name / "provenance.json"
            if not receipt.is_file():
                continue  # Preparation will create or restore the managed link.
            row = json.loads(receipt.read_text(encoding="utf-8"))
            if not row.get("cache_key") or not row.get("payload_root"):
                continue  # Legacy migration is handled by the resource worker.
            payload = Path(row["payload_root"])
            if not payload.is_dir():
                raise ValueError(f"Prepared resource payload is missing: {payload}. Restore the recorded snapshot before resuming.")
            alias = receipt.parent / "data"
            if not alias.is_symlink() or alias.resolve() != payload.resolve():
                raise ValueError(
                    f"Managed resource link is missing or changed: {alias}. With the workflow stopped, "
                    f"move only {receipt} aside and rerun the same command. The verified snapshot "
                    "cache will restore its link and receipt without a download.")

    def _samples(self, supplied):
        # A direct isolate=... command line overrides the saved batch list.
        rows, source = [], self.base
        if supplied.get("isolate"):
            rows = [{"isolate_dir": supplied["isolate"]}]
        elif supplied.get("samples"):
            sheet = Path(_path(supplied["samples"], self.root))
            rows, source = _sheet(sheet), sheet.parent
        else:
            entries = supplied.get("isolates", [])
            if isinstance(entries, str):
                entries = [entries]
            if isinstance(entries, dict):
                rows = [{"sample_id": key, **(value if isinstance(value, dict) else {"isolate_dir": value})}
                        for key, value in entries.items()]
            else:
                rows = [value if isinstance(value, dict) else {"isolate_dir": value} for value in entries]
        if not rows:
            return {}  # The explicit `setup` target does not need sample inputs.
        samples, ids = {}, set()
        for row in rows:
            row = {key: value for key, value in row.items() if value is not None and value != ""}
            directory = row.get("isolate_dir") or row.get("isolate") or row.get("sample_id")
            if not directory:
                raise ValueError("Sample requires isolate_dir or sample_id")
            isolate = _path(directory, source)
            sid = str(row.get("sample_id") or Path(isolate).name)
            if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", sid) or sid in {".", ".."}:
                raise ValueError("Invalid sample_id: " + sid)
            if sid.casefold() in ids or isolate in samples:
                raise ValueError("Duplicate sample ID or isolate directory")
            if any(char in isolate for char in "{}\n\r"):
                raise ValueError("Isolate path cannot contain braces or newlines")
            sample = dict(row, sample_id=sid, isolate_dir=isolate,
                          safety_dir=str(Path(isolate) / "safety"))
            for key, suffix in (("assembly", "fna"), ("proteins", "faa"), ("gff", "gff"), ("genbank", "gbk")):
                path = _path(row[key], source) if row.get(key) else str(Path(isolate) / "annotation" / f"{sid}.{suffix}")
                if key == "genbank" and not Path(path).is_file() and not row.get(key):
                    continue
                if not Path(path).is_file():
                    raise ValueError(f"{sid}: missing {key}: {path}")
                sample[key] = path
            sample.setdefault("input_mode", "auto")
            if sample.get("hybas_dir"):
                sample["hybas_dir"] = _path(sample["hybas_dir"], source)
            sample.setdefault("phenotypes", [])
            sample = enrich_sample(sample)
            samples[isolate] = sample
            ids.add(sid.casefold())
        return samples

    def _resource_specs(self):
        specs = {}
        def request(name, targets, dependencies=(), options=None, external=None):
            owner = PROVIDER_OWNER[name]
            config = self.config["deep"] if owner == "deepsafe" else self.config["modules"][owner]
            row = specs.setdefault(name, {"id": name, "provider": name,
                "recipe": RECIPES[owner], "targets": [], "dependencies": [],
                "module_config": copy.deepcopy(config), "options": {},
                "threads": min(config.get("threads", self.threads), self.threads),
                "memory_mb": min(config.get("memory_mb", 4000), self.memory_mb)})
            row["targets"] = sorted(set(row["targets"]) | set(targets))
            row["dependencies"] = sorted(set(row["dependencies"]) | set(dependencies))
            row["options"].update(options or {})
            if external:
                row.update(mode="external", **external)
            elif self.resource_policy == "manual" and name != "minced":
                raise ValueError("Manual resources require configured paths for " + name)
            return name
        for module in ALL_MODULES:
            if module == "pathogenfinder2":
                continue
            if ((not self.samples and self.config["modules"][module]["enabled"])
                    or any(self.module_enabled(isolate, module) for isolate in self.samples)):
                fields = self.config["modules"][module]
                external = {"module_updates": {module: {key: value for key, value in fields.items()
                            if key in PATH_KEYS or key == "database_version"}}} if fields.get("database") else None
                request(PROVIDERS[module], [module], external=external)
        if self.extend_enabled and ((not self.samples and self.config["modules"]["pathogenfinder2"]["enabled"])
                or any(self.module_enabled(i, "pathogenfinder2") for i in self.samples)):
            fields = self.config["modules"]["pathogenfinder2"]
            external = ({"module_updates": {"pathogenfinder2": fields}}
                        if fields.get("prot_t5_path") and fields.get("weights_model") else None)
            request("prott5", ["pathogenfinder2"], options={"external_modules": {"pathogenfinder2": fields}}, external=external)
        if self.deep_enabled:
            deep = self.config["deep"]
            external = {"deep_updates": {"prosite_snapshot": deep["prosite_snapshot"]}} if deep.get("prosite_snapshot") else None
            request("prosite", ["deepsafe"], external=external)
            dependencies, external_modules = [], {}
            for module in ("amrfinder", "rgi", "vfdb", "antismash"):
                fields = self.config["modules"][module]
                required = (not deep.get("domain_database")) if module == "antismash" else not deep.get("reference_fasta")
                if not required:
                    continue
                if fields.get("database"):
                    external_modules[module] = fields
                else:
                    dependencies.append(request(PROVIDERS[module], ["deepsafe"]))
            options = {"external_modules": external_modules, "domain_database": deep.get("domain_database")}
            if bool(deep.get("reference_fasta")) != bool(deep.get("reference_metadata")):
                raise ValueError("Deep reference FASTA and metadata must be configured together")
            if deep.get("reference_fasta"):
                options.update(preserve_references=True, reference_fasta=deep["reference_fasta"],
                               reference_metadata=deep["reference_metadata"])
            required_fields = ["reference_fasta", "reference_metadata", "domain_database", "profiles_dir"]
            if deep.get("motifs", True):
                required_fields.append("motif_rules")
            if deep.get("family_placement"):
                required_fields.append("placement_registry")
            external = ({"deep_updates": {key: deep[key] for key in required_fields}}
                        if all(deep.get(key) for key in required_fields) else None)
            request("deep_catalogue", ["deepsafe"], dependencies, options, external=external)
        return specs

    def resource_specs(self):
        return copy.deepcopy(self._specs)

    def resource_receipt(self, name):
        return str(Path(self.resource_root) / name / "provenance.json")

    def resource_inputs(self, module):
        return [self.resource_receipt(name) for name, spec in self._specs.items() if module in spec["targets"]]

    def env(self, module):
        return str(self.root / "envs" / (RECIPES.get(module, module) + ".yml"))

    def env_for_resource(self, name):
        return self.env(PROVIDER_OWNER[name])

    def needed_recipes(self):
        result = {RECIPES[module] for module in ALL_MODULES
                  if ((not self.samples and self.config["modules"][module]["enabled"])
                      or any(self.module_enabled(i, module) for i in self.samples))}
        result.update(spec["recipe"] for spec in self._specs.values())
        if self.deep_enabled:
            result.add("deepsafe")
        return sorted(result)

    def environment_receipts(self):
        return [str(Path(self.resource_root) / "environments" / (name + ".json")) for name in self.needed_recipes()]

    def module_environment_receipts(self, module):
        recipe = RECIPES[module]
        return ([str(Path(self.resource_root) / "environments" / (recipe + ".json"))]
                if recipe in self.needed_recipes() else [])

    def module_config_path(self, isolate, module):
        if module not in ALL_MODULES and module != "deepsafe":
            raise ValueError("Unknown module configuration: " + str(module))
        return str(Path(isolate) / "safety/metadata/configs" / (module + ".json"))

    def module_document(self, isolate, module):
        """Configuration consumed by one execution stage, without other policy.

        The complete configuration remains in metadata/config.json for reports.
        Per-module documents are independently produced: changing downstream
        admission policy must not invalidate completed upstream native work.
        """
        result = {"version": self.config["version"], "engine": self.config["engine"],
                  "modules": {}, "module_decisions": {}}
        if module == "deepsafe":
            result["deep"] = copy.deepcopy(self.config["deep"])
        else:
            result["modules"][module] = self.module_config(isolate, module)
            result["module_decisions"][module] = copy.deepcopy(self.decisions[isolate][module])
            if module == "pathogenfinder2":
                result["execution"] = copy.deepcopy(self.config["execution"])
        return result

    def deep_config(self, isolate):
        # DeepSafe currently has one global configuration; keep the isolate
        # argument for the same resource interface used by module rules.
        if str(isolate) not in self.samples:
            raise ValueError("Unknown isolate: " + str(isolate))
        return copy.deepcopy(self.config["deep"])

    def setup_target(self, isolate):
        # Every requested isolate's input/setup check must finish before any
        # biological job begins. This is a DAG dependency, not a Python barrier.
        return [str(Path(item) / "safety/metadata/setup.ready.json") for item in self.samples]

    def sample_file(self, isolate, key):
        return self.samples[str(isolate)][key]

    def module_enabled(self, isolate, module):
        return self.decisions[str(isolate)][module]["action"] == "run"

    def module_config(self, isolate, module):
        fields = copy.deepcopy(self.config["modules"][module])
        # Command construction happens when the DAG is built; resource outputs
        # have stable paths, while version/resolved-path details come in receipts.
        name = PROVIDERS.get(module)
        if not fields.get("database") and name in self._specs:
            data = Path(self.resource_root) / name / "data"
            suffixes = {"amrfinder": "database/latest", "rgi": "card.json", "vfdb": "vfdb_setA.dmnd",
                        "bacmet": "bacmet_exp.dmnd", "genomad": "genomad_db", "mlst": "db",
                        "crisprcasfinder": "source"}
            fields["database"] = str(data / suffixes.get(module, ""))
        if module == "icescreen" and self.module_enabled(isolate, module):
            fields["phylum"] = "bacillota"
        return fields

    def sample_config(self, isolate):
        cfg = copy.deepcopy(self.config)
        cfg["modules"] = {module: self.module_config(isolate, module) for module in ALL_MODULES}
        cfg["module_decisions"] = self.decisions[isolate]
        cfg["outdir"] = str(Path(isolate) / "safety")
        return cfg

    def evidence_targets(self, isolate):
        return [str(Path(isolate) / "safety/evidence" / (name + ".json"))
                for name in ALL_MODULES if self.decisions[isolate][name]["action"] != "disabled"]

    def screening_report_targets(self):
        if not self.samples:
            raise ValueError("No samples selected. Use the setup target to install only, or supply isolates/samples.")
        return [str(Path(isolate) / "safety" / relative) for isolate in self.samples
                for relative in ("report.html", "tables/amr-review.tsv")]

    def followup_report_targets(self):
        if not self.samples:
            raise ValueError("No samples selected. Use the setup target to install only, or supply isolates/samples.")
        if not self.deep_enabled:
            return []
        return [str(Path(isolate) / "safety/deepsafe-report" / relative) for isolate in self.samples
                for relative in ("report.html", "tables/amr-review.tsv")]

    def report_targets(self):
        return self.screening_report_targets() + self.followup_report_targets()
